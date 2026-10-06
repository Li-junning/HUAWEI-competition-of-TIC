"""Real DOCX bytes exercise extraction, indexing, provenance and failures."""

import base64
import io
import sys
import zipfile

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.application import create_app
from app.config import Settings
from app.knowledge import KnowledgeBase, KnowledgeError
from app.knowledge_schemas import KnowledgeImport
from app.storage import Storage


def word_bytes(build=None):
    document = Document()
    if build:
        build(document)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def word_body(data, filename="资料.docx"):
    return {"title": "Word 参考资料", "filename": filename,
            "file_base64": base64.b64encode(data).decode(), "tags": ["技术"]}


def research_document(document):
    document.add_paragraph("研究记录 🧪")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("研究主体", "研究结论"), ("海川公司研发中心", "位于南京。")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    document.add_paragraph("结论仅适用于本研究所列条件。")


def test_docx_api_import_search_source_duplicate_and_restart(tmp_path):
    database = tmp_path / "word.db"
    data = word_bytes(research_document)
    expected = "研究记录 🧪\n研究主体\t研究结论\n海川公司研发中心\t位于南京。\n结论仅适用于本研究所列条件。"
    with TestClient(create_app(Settings(database_path=database))) as client:
        response = client.post("/api/knowledge/documents", json=word_body(data, "中文研究.DOCX"))
        assert response.status_code == 201
        doc = response.json()
        assert doc["filename"] == "中文研究.DOCX"
        assert doc["char_count"] == len(expected)
        assert doc["page_count"] == 1 and doc["chunk_count"] > 0
        source = client.get(f"/api/knowledge/documents/{doc['document_id']}").json()
        assert source["pages"] == [expected]
        hits = client.post("/api/knowledge/search", json={"query": "研发中心 南京", "tag": "技术"}).json()["items"]
        assert hits
        for hit in hits:
            assert hit["document_id"] == doc["document_id"]
            assert hit["page"] is None
            assert expected[hit["char_start"]:hit["char_end"]] == hit["excerpt"]
            assert hit["content_hash"] == doc["content_hash"]
        duplicate = client.post("/api/knowledge/documents", json=word_body(data, "另一名称.docx"))
        assert duplicate.status_code == 409
        assert duplicate.json()["error"]["code"] == "KB_DUPLICATE"
    with TestClient(create_app(Settings(database_path=database))) as client:
        assert client.get(f"/api/knowledge/documents/{doc['document_id']}").json()["pages"] == [expected]
        assert client.post("/api/knowledge/search", json={"query": "南京"}).json()["items"]


def test_docx_tables_keep_nested_order_and_merged_content_once(tmp_path):
    def build(document):
        document.add_paragraph("表格之前")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).merge(table.cell(0, 1)).text = "合并标题"
        cell = table.cell(1, 0)
        cell.text = "单元格开头"
        nested = cell.add_table(rows=1, cols=2)
        nested.cell(0, 0).text = "嵌套甲"
        nested.cell(0, 1).text = "嵌套乙"
        cell.add_paragraph("嵌套之后")
        table.cell(1, 1).text = "第二列"
        document.add_paragraph("表格之后")

    storage = Storage(tmp_path / "tables.db")
    try:
        library = KnowledgeBase(storage)
        doc = library.import_document(KnowledgeImport(**word_body(word_bytes(build))))
        text = library.get_document(doc.document_id)["pages"][0]
        assert text.count("合并标题") == 1
        ordered = ["表格之前", "合并标题", "单元格开头", "嵌套甲", "嵌套乙", "嵌套之后", "第二列", "表格之后"]
        assert [text.index(value) for value in ordered] == sorted(text.index(value) for value in ordered)
        assert "嵌套甲\t嵌套乙" in text
        assert library.search("嵌套甲")[0].excerpt in text
    finally:
        storage.close()


def test_docx_preserves_hyperlink_text_without_fetching_the_target(tmp_path):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE

    def build(document):
        paragraph = document.add_paragraph("参考 ")
        hyperlink = OxmlElement("w:hyperlink")
        hyperlink.set(qn("r:id"), document.part.relate_to("https://example.invalid/no-network", RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
        run, text = OxmlElement("w:r"), OxmlElement("w:t")
        text.text = "公开研究结论"
        run.append(text)
        hyperlink.append(run)
        paragraph._p.append(hyperlink)

    with TestClient(create_app(Settings(database_path=tmp_path / "hyperlink.db"))) as client:
        response = client.post("/api/knowledge/documents", json=word_body(word_bytes(build)))
        assert response.status_code == 201
        source = client.get("/api/knowledge/documents/" + response.json()["document_id"]).json()
        assert source["pages"] == ["参考 公开研究结论"]


def archive_bytes(parts):
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        for name, content in parts.items():
            archive.writestr(name, content)
    return stream.getvalue()


@pytest.mark.parametrize("data,filename,status,code", [
    (b"not a Word document", "损坏.docx", 400, "KB_WORD_INVALID"),
    (archive_bytes({"other.xml": b"text"}), "不是Word.docx", 400, "KB_WORD_INVALID"),
    (archive_bytes({"word/document.xml": b"<broken>"}), "XML损坏.docx", 400, "KB_WORD_INVALID"),
    (word_bytes(), "空白.docx", 400, "KB_TEXT_EMPTY"),
    (b"old Word bytes", "旧格式.doc", 400, "KB_WORD_LEGACY"),
    (b"x" * (2 * 1024 * 1024 + 1), "超大.docx", 413, "KB_FILE_TOO_LARGE"),
    (word_bytes(lambda document: document.add_paragraph("字" * 100_001)), "正文超长.docx", 413, "KB_TEXT_LIMIT"),
    (archive_bytes({"word/document.xml": b"x" * (8 * 1024 * 1024 + 1)}), "压缩正文过大.docx", 413, "KB_WORD_LIMIT"),
    (archive_bytes({"word/document.xml": b"<document/>", "word/media/image": b"x" * (32 * 1024 * 1024)}), "压缩包过大.docx", 413, "KB_WORD_LIMIT"),
], ids=["corrupt", "wrong-package", "broken-xml", "empty", "legacy-doc", "file-too-large", "text-too-long", "expanded-body-too-large", "expanded-package-too-large"])
def test_docx_api_rejects_invalid_or_oversized_files_without_saving(tmp_path, data, filename, status, code):
    with TestClient(create_app(Settings(database_path=tmp_path / "invalid-word.db"))) as client:
        response = client.post("/api/knowledge/documents", json=word_body(data, filename))
        assert response.status_code == status
        assert response.json()["error"]["code"] == code
        assert client.get("/api/knowledge/status").json()["document_count"] == 0


def test_missing_word_dependency_returns_actionable_error(tmp_path, monkeypatch):
    data = word_bytes(lambda document: document.add_paragraph("文档正文。"))
    monkeypatch.setitem(sys.modules, "docx", None)
    storage = Storage(tmp_path / "missing-dependency.db")
    try:
        # Dependency availability is checked in the isolated parser now.
        from app.knowledge import _read_docx
        with pytest.raises(KnowledgeError) as caught:
            _read_docx(data)
        assert caught.value.code == "KB_WORD_UNAVAILABLE"
        assert caught.value.status == 503
    finally:
        storage.close()
